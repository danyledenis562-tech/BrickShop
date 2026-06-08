# -*- coding: utf-8 -*-
"""Generate Brick Shop diploma defense Q&A as Word (.docx)."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTPUT = "Brick-Shop-Defense-QA.docx"

# (category, question, answer)
QA = [
    # --- Загальні ---
    ("Загальні питання", "Що таке Brick Shop?", "Веб-застосунок інтернет-магазину LEGO-наборів: каталог, кошик, checkout, оплата, профіль, адмін-панель."),
    ("Загальні питання", "Яка мета дипломної роботи?", "Спроєктувати й реалізувати повноцінний e-commerce застосунок з реальною бізнес-логікою та інтеграціями."),
    ("Загальні питання", "Які основні модулі системи?", "Каталог, товар, кошик, checkout, оплата, профіль, обране, адмін-панель, локалізація."),
    ("Загальні питання", "Хто користувачі системи?", "Покупець (гість або зареєстрований), адміністратор магазину."),
    ("Загальні питання", "Чим проєкт відрізняється від простого CRUD?", "Є checkout, оплата, доставка, бонуси, промокоди, модерація відгуків, інтеграції з API."),
    ("Загальні питання", "Яка практична цінність?", "Готовий до запуску магазин: локально (XAMPP) і на сервері (Docker/Render)."),
    ("Загальні питання", "Скільки мов підтримує інтерфейс?", "4: українська (за замовчуванням), англійська, польська, російська."),
    ("Загальні питання", "Чи можна купити без реєстрації?", "Так, реалізовано гостьовий checkout з email і signed URL на сторінку подяки."),
    ("Загальні питання", "Що демонструвати на захисті?", "Каталог → товар → кошик → checkout → оплата → профіль → адмінка → перемикання мови."),
    ("Загальні питання", "Демо-акаунти для показу?", "Admin: admin@brickshop.test / Admin123!; User: user@brickshop.test / User123!"),

    # --- Технології ---
    ("Технології", "Який backend-фреймворк?", "Laravel 10 на PHP 8.1+."),
    ("Технології", "Який frontend?", "Blade-шаблони, Vite, Tailwind CSS, Alpine.js, окремі JS-модулі."),
    ("Технології", "Яка база даних локально?", "MySQL через XAMPP."),
    ("Технології", "Яка БД на production?", "PostgreSQL на Render."),
    ("Технології", "Чому Laravel, а не Node.js?", "Швидша розробка: auth, міграції, валідація, mail, middleware «з коробки»."),
    ("Технології", "Чому Blade, а не React SPA?", "SEO, простіші сесії/CSRF, менше JS-складності для MVP."),
    ("Технології", "Що таке Vite у проєкті?", "Збірник фронтенду: CSS/JS, hot reload у dev, оптимізований build для prod."),
    ("Технології", "Що таке Tailwind CSS?", "Utility-first CSS-фреймворк для швидкої верстки без великої кількості власного CSS."),
    ("Технології", "Для чого Alpine.js?", "Легка інтерактивність у Blade (модалки, dropdown) без важкого SPA."),
    ("Технології", "Як реалізована автентифікація?", "Laravel Breeze — session-based login/register."),
    ("Технології", "Чи є social login?", "Так, Google OAuth через Laravel Socialite."),
    ("Технології", "Що таке Laravel Sanctum у проєкті?", "Підключений мінімально; основний доступ через web-сесії, не REST API."),
    ("Технології", "Яка архітектурна модель?", "MVC + Service Layer: контролери тонкі, логіка в сервісах."),
    ("Технології", "Навіщо Service Layer?", "Винести бізнес-логіку з контролерів: cart, checkout, pricing, payments."),
    ("Технології", "Що таке middleware?", "Проміжний шар: locale, auth, admin, CSRF, throttle до контролера."),
    ("Технології", "Що таке Form Request?", "Клас валідації вхідних даних (CheckoutRequest, CartRequest тощо)."),
    ("Технології", "Що таке Enum у проєкті?", "OrderStatus: new, paid, processing, shipped, canceled — type-safe статуси."),
    ("Технології", "Що таке route model binding?", "Маршрут /product/{product:slug} — Laravel сам знаходить товар по slug."),

    # --- Архітектура / файли ---
    ("Архітектура", "Де логіка каталогу?", "CatalogController — фільтри, сортування, пошук, пагінація."),
    ("Архітектура", "Де логіка кошика?", "CartService (session) + CartController (HTTP)."),
    ("Архітектура", "Де створюється замовлення?", "OrderCreator у DB-транзакції."),
    ("Архітектура", "Де рахується ціна checkout?", "CheckoutPricingService: subtotal, promo, shipping, bonus."),
    ("Архітектура", "Де оплата LiqPay?", "LiqPayService + LiqPayCallbackController."),
    ("Архітектура", "Де Nova Poshta?", "NovaPoshtaClient + ShippingController (JSON endpoints)."),
    ("Архітектура", "Як захищена адмінка?", "Middleware auth + admin (EnsureAdmin, role === 'admin')."),
    ("Архітектура", "Скільки контролерів?", "32: публічні, auth, payments, admin."),
    ("Архітектура", "Де переклади?", "lang/uk|en|pl|ru/messages.php + SetLocale middleware."),
    ("Архітектура", "Де конфіг магазину?", "config/shop.php + змінні .env."),

    # --- База даних ---
    ("База даних", "Основні таблиці?", "users, products, categories, product_images, orders, order_items, promo_codes, reviews, bonus_transactions."),
    ("База даних", "Зв'язок users ↔ orders?", "Один користувач — багато замовлень; user_id nullable для гостей."),
    ("База даних", "Зв'язок categories ↔ products?", "Одна категорія — багато товарів (category_id)."),
    ("База даних", "Навіщо product_images окремо?", "Кілька фото на товар, прапорець is_main для обкладинки."),
    ("База даних", "Навіщо order_items?", "Snapshot позицій: qty, price, total на момент покупки."),
    ("База даних", "Чому snapshot ціни?", "Ціна в products може змінитись — історія замовлень лишається коректною."),
    ("База даних", "Що в orders для гостя?", "user_id = null, guest_email заповнений."),
    ("База даних", "Як працює обране?", "Pivot-таблиця favorites: users ↔ products (many-to-many)."),
    ("База даних", "Як recently viewed?", "Auth: pivot recently_viewed; гість: session (max 8)."),
    ("База даних", "Що в promo_codes?", "code, type (percent/fixed), value, дати, usage_limit, times_used."),
    ("База даних", "Що в reviews?", "rating 1–5, comment, approved (модерація)."),
    ("База даних", "Що в bonus_transactions?", "type earn/spend, amount, order_id — журнал бонусів."),
    ("База даних", "Навіщо міграції?", "Версіонування схеми БД, однакова структура на dev і prod."),
    ("База даних", "Навіщо seeders?", "Демо-дані: категорії, товари, користувачі, замовлення для тестів."),
    ("База даних", "Що таке slug у products?", "URL-friendly унікальний ідентифікатор товару (/product/{slug})."),
    ("База даних", "Що таке set_number?", "Номер LEGO-набору (напр. 75192) для пошуку."),

    # --- Каталог і пошук ---
    ("Каталог і пошук", "Які фільтри в каталозі?", "Категорія, ціна min/max, складність, в наявності."),
    ("Каталог і пошук", "Які варіанти сортування?", "За ціною ↑↓, популярністю, новизною (default)."),
    ("Каталог і пошук", "Скільки товарів на сторінку?", "12 з пагінацією."),
    ("Каталог і пошук", "Як працює пошук?", "LIKE по name, series, brand, description, set_number (case-insensitive)."),
    ("Каталог і пошук", "Пошук по номеру набору?", "Окрема нормалізація цифр: 75-192 знаходить 75192."),
    ("Каталог і пошук", "Що таке search suggestions?", "AJAX autocomplete: GET /search-suggestions, max 5 результатів."),
    ("Каталог і пошук", "Де JS для підказок?", "resources/js/ui/searchSuggestions.js — debounce + fetch."),
    ("Каталог і пошук", "Чи кешується каталог?", "Список категорій кешується на 15 хвилин."),
    ("Каталог і пошук", "Які товари показуються?", "Лише is_active = true."),
    ("Каталог і пошук", "Що таке N+1 і як уникнути?", "Зайві запити в циклі; eager loading: with('coverImage', 'category')."),

    # --- Кошик ---
    ("Кошик", "Де зберігається кошик?", "У PHP session, ключ cart (CartService)."),
    ("Кошик", "Чому не в БД?", "Простіше для MVP; не потрібна авторизація для кошика."),
    ("Кошик", "Структура кошика?", "Масив {product_id, quantity}, qty обмежено 1–99."),
    ("Кошик", "Що при перевищенні stock?", "Кількість обмежується наявністю на складі."),
    ("Кошик", "Що після checkout?", "Кошик очищається (CartService::clear)."),
    ("Кошик", "Abandoned cart?", "CartReminderService для logged-in + email-нагадування (cron)."),
    ("Кошик", "Чи синхронізується між пристроями?", "Ні — обмеження session cart (можливе покращення)."),

    # --- Checkout і замовлення ---
    ("Checkout", "Кроки checkout?", "Кошик → форма → pricing → OrderCreator → email → оплата/thanks."),
    ("Checkout", "Хто може оформити?", "Авторизований користувач або гість."),
    ("Checkout", "Як рахується total?", "subtotal − promo + shipping − bonus_spent."),
    ("Checkout", "Вартість доставки?", "Нова Пошта 100 грн, кур'єр 250, Укрпошта 50 (config/shop.php)."),
    ("Checkout", "Як працює промокод?", "isValid() перевіряє дати/ліміт; applyDiscount() — percent або fixed."),
    ("Checkout", "Як працюють бонуси?", "Списання з bonus_balance; нарахування за SHOP_BONUS_EARN_RATE (1 бонус / 10 грн)."),
    ("Checkout", "Що при нестачі товару?", "OutOfStockException → redirect з повідомленням."),
    ("Checkout", "Як захищено stock?", "DB-транзакція + where('stock','>=',qty)->decrement()."),
    ("Checkout", "Статуси замовлення?", "new → paid → processing → shipped або canceled."),
    ("Checkout", "Коли paid?", "Після успішного LiqPay callback."),
    ("Checkout", "Email після замовлення?", "OrderPlacedMail — dispatch afterResponse (не блокує відповідь)."),
    ("Checkout", "Як гість бачить thanks?", "Signed URL — криптографічний підпис Laravel."),
    ("Checkout", "IDOR на thanks?", "Перевірка signature або ownership — чужий order недоступний."),
    ("Checkout", "Throttle checkout?", "10 запитів/хв — захист від spam-замовлень."),

    # --- Оплата ---
    ("Оплата", "Яка платіжна система?", "LiqPay (production) + test webhook для dev."),
    ("Оплата", "Як працює LiqPay flow?", "Форма → оплата на LiqPay → POST callback → verify signature → paid."),
    ("Оплата", "Як перевіряється callback?", "SHA1 signature, hash_equals — без валідного підпису відхиляється."),
    ("Оплата", "Чому CSRF вимкнено для callback?", "Зовнішній POST без CSRF-токена; захист — підпис LiqPay."),
    ("Оплата", "Повторний callback?", "Ідемпотентність: повторна оплата не дублює зміни."),
    ("Оплата", "Тестова оплата в dev?", "TestPaymentController + webhook secret у .env."),

    # --- Доставка ---
    ("Доставка", "Які способи доставки?", "Nova Poshta, кур'єр, Укрпошта."),
    ("Доставка", "Як працює Nova Poshta API?", "NovaPoshtaClient: cities, branches, streets — JSON autocomplete."),
    ("Доставка", "Кеш Nova Poshta?", "6–12 годин — менше навантаження на API."),
    ("Доставка", "Fallback без API key?", "NovaPoshtaFallbackDirectory — offline-дані для demo/dev."),
    ("Доставка", "Де endpoints доставки?", "GET /shipping/nova/cities, branches, streets."),

    # --- Профіль і користувач ---
    ("Профіль", "Що в особистому кабінеті?", "Замовлення, обране, бонуси, нещодавно переглянуті, редагування профілю."),
    ("Профіль", "Завантаження аватара?", "ProfileController — збереження в storage, PublicMedia для URL."),
    ("Профіль", "Скасування замовлення?", "Якщо статус дозволяє — canceled + canceled_at."),
    ("Профіль", "Відстеження посилки?", "tracking_number, tracking_url; email OrderTrackingMail з адмінки."),
    ("Профіль", "Обране без входу?", "Редірект на login — favorites потребує auth."),

    # --- Відгуки ---
    ("Відгуки", "Хто може залишити відгук?", "Лише авторизований користувач."),
    ("Відгуки", "Модерація?", "approved = false за замовчуванням; адмін схвалює в /admin/reviews."),
    ("Відгуки", "Rate limit відгуків?", "5/хв — throttle middleware."),
    ("Відгуки", "Що показується на товарі?", "Лише approved відгуки + середній рейтинг."),

    # --- Адмін-панель ---
    ("Адмін-панель", "Що може адмін?", "CRUD товарів/категорій, замовлення, промокоди, відгуки, користувачі, налаштування."),
    ("Адмін-панель", "Як увійти в адмінку?", "Користувач з role=admin → /admin."),
    ("Адмін-панель", "Звичайний user в /admin?", "403 Forbidden через EnsureAdmin."),
    ("Адмін-панель", "Зміна статусу замовлення?", "Admin\\OrderController + OrderStatusRequest."),
    ("Адмін-панель", "TTN (трекінг)?", "Адмін додає tracking_number → email користувачу/гостю."),

    # --- Безпека ---
    ("Безпека", "Захист від CSRF?", "VerifyCsrfToken на всіх web-формах + @csrf у Blade."),
    ("Безпека", "Захист від XSS?", "Blade {{ }} auto-escaping; не {!! !!} для user input."),
    ("Безпека", "Захист від SQL injection?", "Eloquent ORM, parameterized whereRaw."),
    ("Безпека", "Захист паролів?", "Hash::make, cast hashed у моделі User."),
    ("Безпека", "Brute-force login?", "Throttle на auth routes."),
    ("Безпека", "Mass assignment?", "$fillable на моделях — лише дозволені поля."),
    ("Безпека", "Signed URLs?", "checkout.thanks для гостей, email verification."),
    ("Безпека", "Розділення ролей?", "users.role: admin / user + middleware."),
    ("Безпека", "Path traversal у media?", "MediaController — лише public disk, 404 для ../."),
    ("Безпека", ".env доступний з web?", "Ні — 404/403 від веб-сервера."),
    ("Безпека", "Google OAuth безпека?", "stateless() у production; google_id у users."),

    # --- Локалізація ---
    ("Локалізація", "Як перемикається мова?", "GET /locale/{locale} → LocaleController → session + cookie."),
    ("Локалізація", "Де зберігається вибір?", "Session і cookie locale."),
    ("Локалізація", "Мова за замовчуванням?", "uk (config/app.php)."),
    ("Локалізація", "Fallback мова?", "en — якщо переклад відсутній."),

    # --- Тестування ---
    ("Тестування", "Який test framework?", "PHPUnit 10, Feature tests з RefreshDatabase."),
    ("Тестування", "Скільки тестових файлів?", "18 файлів, ~50+ методів у tests/Feature."),
    ("Тестування", "Що тестує CatalogTest?", "Каталог, пошук, set_number, сторінка товару."),
    ("Тестування", "Що тестує CheckoutTest?", "Створення order, очищення кошика, mail, thanks access."),
    ("Тестування", "Що тестує GuestCheckoutTest?", "Гостьовий order + signed thanks URL."),
    ("Тестування", "Factories?", "ProductFactory, UserFactory, OrderFactory — тестові дані."),
    ("Тестування", "Mail у тестах?", "Mail::fake() — перевірка без реальної відправки."),
    ("Тестування", "Статичний аналіз?", "PHPStan/Larastan level 0: composer analyse."),
    ("Тестування", "Чому Feature, а не Unit?", "Перевіряємо повний HTTP-flow, ближче до реального користувача."),
    ("Тестування", "Є ручний тест-план?", "Так, docs/Brick-Shop-Test-Plan.docx."),

    # --- Інтеграції ---
    ("Інтеграції", "Зовнішні API?", "Nova Poshta, LiqPay, Google OAuth, Cloudinary (опційно)."),
    ("Інтеграції", "Чи є REST API для клієнтів?", "Мінімально: GET /api/user (Sanctum); основа — web routes."),
    ("Інтеграції", "Email-сервіси?", "SMTP/log/Mailgun/Postmark через config mail."),
    ("Інтеграції", "Cloudinary для чого?", "CDN для product images: shop:sync-product-images-to-cloudinary."),
    ("Інтеграції", "Health check?", "GET /health — для Render/deployment monitoring."),

    # --- Деплой ---
    ("Деплой", "Як запустити локально?", "composer install, npm install, migrate --seed, npm run dev, artisan serve."),
    ("Деплой", "XAMPP особливості?", "APP_URL з /public; MySQL на 127.0.0.1."),
    ("Деплой", "Production deploy?", "Docker + Render (render.yaml, start.sh)."),
    ("Деплой", "Що робить start.sh?", "migrate, optional seed, mirror images, cache config/routes/views."),
    ("Деплой", "Production cache?", "config:cache, route:cache, view:cache."),
    ("Деплой", "Session на Render?", "SESSION_DRIVER=database."),
    ("Деплой", "Зображення на проді?", "mirror-product-images або Cloudinary sync."),
    ("Деплой", "storage:link навіщо?", "Симлінк public/storage → storage/app/public для фото."),

    # --- Artisan / команди ---
    ("Команди", "shop:mirror-product-images?", "Завантажує зовнішні URL фото в локальне сховище."),
    ("Команди", "shop:sync-product-images-to-cloudinary?", "Синхронізує фото в Cloudinary CDN."),
    ("Команди", "shop:send-abandoned-cart-reminders?", "Розсилка нагадувань про покинутий кошик (cron)."),
    ("Команди", "php artisan migrate?", "Застосовує міграції без скидання даних."),
    ("Команди", "migrate:fresh небезпечний?", "Так — повністю очищає всі таблиці."),

    # --- Обмеження і перспективи ---
    ("Обмеження", "Головні обмеження MVP?", "Session cart, LIKE-пошук, мінімальний API, без WebSocket stock."),
    ("Обмеження", "Кошик між пристроями?", "Не синхронізується — потрібен DB cart."),
    ("Обмеження", "Масштабування пошуку?", "LIKE не для великих каталогів — Meilisearch/Elasticsearch."),
    ("Обмеження", "Горизонтальне масштабування?", "Session cart потребує Redis/database sessions."),
    ("Перспективи", "Що покращити далі?", "DB cart, full-text search, Redis queue, mobile API, аналітика."),
    ("Перспективи", "Мобільний додаток?", "Потрібен повноцінний REST API + Sanctum tokens."),
    ("Перспективи", "Real-time stock?", "WebSocket або polling при checkout."),

    # --- Швидкі «пастки» ---
    ("Швидкі відповіді", "Де кошик?", "PHP session, ключ cart."),
    ("Швидкі відповіді", "Де ціна в замовленні?", "order_items.price — snapshot."),
    ("Швидкі відповіді", "Скільки мов?", "4: uk, en, pl, ru."),
    ("Швидкі відповіді", "Яка БД локально / prod?", "MySQL / PostgreSQL."),
    ("Швидкі відповіді", "Як гість бачить thanks?", "Signed URL."),
    ("Швидкі відповіді", "Як захищено LiqPay?", "SHA1 signature verify."),
    ("Швидкі відповіді", "Ролі користувачів?", "admin і user."),
    ("Швидкі відповіді", "Статус після оплати?", "paid."),
    ("Швидкі відповіді", "Відгуки на сайті?", "Лише approved."),
    ("Швидкі відповіді", "Товарів на сторінку?", "12."),
    ("Швидкі відповіді", "Бонуси: скільки за 100 грн?", "10 бонусів (rate 10)."),
    ("Швидкі відповіді", "Доставка НП?", "100 грн."),
    ("Швидкі відповіді", "Фреймворк?", "Laravel 10."),
    ("Швидкі відповіді", "CSS?", "Tailwind CSS."),
    ("Швидкі відповіді", "Збірник JS?", "Vite 5."),
]

HEADERS = ["№", "Категорія", "Питання", "Коротка відповідь"]


def set_cell_shading(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)

    title = doc.add_heading("Brick Shop — питання та відповіді на захист", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(
        "Дипломний проєкт: lego-shop3 (Laravel 10 + Blade + Vite)\n"
        f"Кількість питань: {len(QA)}\n"
        "Формат: короткі відповіді для усного захисту"
    )
    run.font.size = Pt(10)

    doc.add_paragraph()

    table = doc.add_table(rows=1, cols=len(HEADERS))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = table.rows[0].cells
    for i, text in enumerate(HEADERS):
        hdr_cells[i].text = text
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(9)
        set_cell_shading(hdr_cells[i], "D9E2F3")

    widths = [Cm(0.8), Cm(2.8), Cm(5.5), Cm(8.5)]
    for i, w in enumerate(widths):
        hdr_cells[i].width = w

    current_category = None
    for idx, (cat, question, answer) in enumerate(QA, start=1):
        cells = table.add_row().cells
        show_cat = cat if cat != current_category else ""
        if cat != current_category:
            current_category = cat

        values = [str(idx), show_cat, question, answer]
        for i, val in enumerate(values):
            cells[i].text = val
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)
                    if i == 2:
                        r.bold = True

    doc.add_page_break()
    doc.add_heading("Сценарій демонстрації (5–7 хв)", level=1)
    demo_steps = [
        "Головна — featured товари, категорії",
        "Каталог — фільтр, сортування, пошук (напр. 75192)",
        "Сторінка товару — галерея, відгук, «в обране»",
        "Кошик → Checkout (гість або user@brickshop.test)",
        "Оплата (тестовий режим) → сторінка подяки",
        "Профіль — замовлення, бонуси",
        "Адмінка (admin@brickshop.test) — статус замовлення, модерація відгуку",
        "Перемикання мови uk ↔ en",
    ]
    for step in demo_steps:
        doc.add_paragraph(step, style="List Number")

    doc.add_paragraph()
    acc = doc.add_paragraph()
    acc.add_run("Демо-акаунти: ").bold = True
    acc.add_run("Admin — admin@brickshop.test / Admin123!  |  User — user@brickshop.test / User123!")

    import os
    script_dir = os.path.dirname(os.path.abspath(__file__))
    full_path = os.path.join(script_dir, OUTPUT)
    doc.save(full_path)
    print(full_path)


if __name__ == "__main__":
    main()
