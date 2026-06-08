# -*- coding: utf-8 -*-
"""Generate Brick Shop manual test plan as Word (.docx)."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTPUT = "Brick-Shop-Test-Plan.docx"

# (category, id, module, name, preconditions, steps, expected, priority)
TESTS = [
    # --- Функціональне ---
    ("Функціональне тестування", "FN-01", "Головна /", "Відображення головної сторінки",
     "Сайт запущено, є товари в БД",
     "1. Відкрити /\n2. Перевірити блоки: банери, featured, новинки, категорії",
     "Сторінка 200 OK; картки товарів з ціною, зображенням і посиланням на /product/{slug}",
     "Високий"),
    ("Функціональне тестування", "FN-02", "Каталог /catalog", "Фільтрація за категорією",
     "Є ≥2 категорії з товарами",
     "1. Відкрити каталог\n2. Обрати категорію\n3. Застосувати фільтр",
     "Список містить лише товари обраної категорії; лічильник результатів коректний",
     "Високий"),
    ("Функціональне тестування", "FN-03", "Каталог /catalog", "Пошук і підказки",
     "Є товар з відомою назвою/артикулом",
     "1. Ввести частину назви в пошук\n2. Перевірити /search-suggestions\n3. Підтвердити пошук",
     "Підказки з’являються; результати містять релевантні товари",
     "Високий"),
    ("Функціональне тестування", "FN-04", "Товар /product/{slug}", "Сторінка товару та галерея",
     "Товар з ≥2 зображеннями",
     "1. Відкрити картку товару\n2. Перемкнути фото в галереї\n3. Перевірити ціну, наявність, опис",
     "Галерея працює; відображаються ціна, old_price (якщо є), кнопка «В кошик»",
     "Високий"),
    ("Функціональне тестування", "FN-05", "Відгуки", "Додавання відгуку (авторизований)",
     "Користувач увійшов, товар існує",
     "1. На сторінці товару обрати рейтинг і текст\n2. Надіслати відгук",
     "Відгук з’являється після модерації або зі статусом pending; гість не може залишити відгук",
     "Середній"),
    ("Функціональне тестування", "FN-06", "Кошик /cart", "Додавання та зміна кількості",
     "Товар у наявності",
     "1. Додати товар у кошик\n2. Збільшити/зменшити qty\n3. Оновити сторінку кошика",
     "Сума перераховується; кількість не перевищує stock",
     "Високий"),
    ("Функціональне тестування", "FN-07", "Кошик /cart", "Видалення позиції",
     "У кошику ≥1 позиція",
     "1. Натиснути видалити\n2. Перевірити порожній кошик",
     "Позиція зникла; підсумок 0; перехід на checkout недоступний або з попередженням",
     "Високий"),
    ("Функціональне тестування", "FN-08", "Оформлення /checkout", "Гостьове замовлення",
     "Кошик не порожній",
     "1. Перейти на checkout без входу\n2. Заповнити ПІБ, телефон, email, доставку\n3. Підтвердити",
     "Створено order; redirect на /checkout/thanks/{id}; лист order-placed (якщо mail налаштовано)",
     "Високий"),
    ("Функціональне тестування", "FN-09", "Оформлення /checkout", "Доставка Нова Пошта",
     "API Nova Poshta доступне",
     "1. Обрати НП\n2. Знайти місто (typeahead)\n3. Обрати відділення/вулицю",
     "Поля заповнюються; адреса зберігається в замовленні у читабельному вигляді",
     "Високий"),
    ("Функціональне тестування", "FN-10", "Оформлення /checkout", "Промокод",
     "Є активний promo-code у БД",
     "1. Ввести валідний код\n2. Ввести прострочений/невалідний",
     "Знижка застосована до підсумку; невалідний — повідомлення про помилку",
     "Середній"),
    ("Функціональне тестування", "FN-11", "Оплата", "LiqPay callback / тестова оплата",
     "LIQPAY_* або test payment увімкнено",
     "1. Оформити замовлення\n2. Пройти оплату (sandbox)\n3. Перевірити статус order",
     "Після callback статус оновлено; повторний callback ідемпотентний",
     "Високий"),
    ("Функціональне тестування", "FN-12", "Авторизація", "Реєстрація та вхід",
     "Унікальний email",
     "1. /register — створити акаунт\n2. /login — увійти\n3. /logout",
     "Сесія створюється; після logout захищені сторінки недоступні",
     "Високий"),
    ("Функціональне тестування", "FN-13", "Авторизація", "Google OAuth",
     "GOOGLE_CLIENT_ID налаштовано",
     "1. Натиснути «Увійти через Google»\n2. Завершити OAuth",
     "Користувач створений/прив’язаний; google_id у БД",
     "Середній"),
    ("Функціональне тестування", "FN-14", "Профіль /profile", "Редагування та замовлення",
     "Користувач з історією замовлень",
     "1. Відкрити профіль\n2. Змінити ім’я/телефон\n3. Скасувати замовлення (якщо дозволено)",
     "Дані збережені; статус замовлення cancelled; бонуси відображаються (якщо увімкнено)",
     "Середній"),
    ("Функціональне тестування", "FN-15", "Обране /favorites", "Додати/прибрати",
     "Користувач авторизований",
     "1. Додати товар в обране з каталогу\n2. Відкрити /favorites\n3. Видалити",
     "Список синхронізований; без входу — редірект на login",
     "Середній"),
    ("Функціональне тестування", "FN-16", "Локалізація", "Перемикач мови",
     "Увімкнено uk/en/ru/pl",
     "1. Перейти /locale/en\n2. Перевірити підписи в header і каталозі",
     "Тексти з lang/*.php; вибір зберігається в сесії",
     "Середній"),
    ("Функціональне тестування", "FN-17", "Адмін /admin", "CRUD товарів і категорій",
     "Користувач role=admin",
     "1. Увійти як admin\n2. Створити категорію і товар з фото\n3. Редагувати ціну\n4. Видалити",
     "Зміни в БД; товар видно на вітрині після publish/active",
     "Високий"),
    ("Функціональне тестування", "FN-18", "Адмін /admin/orders", "Обробка замовлення",
     "Є нове замовлення",
     "1. Відкрити order в адмінці\n2. Змінити статус\n3. Експорт CSV",
     "Статус оновлено; tracking mail (якщо налаштовано); export завантажується",
     "Високий"),
    ("Функціональне тестування", "FN-19", "Розсилка", "Підписка newsletter",
     "—",
     "1. Ввести email у форму footer/widget\n2. Повторно той самий email",
     "Підписка збережена; дубль — валідація/throttle",
     "Низький"),
    ("Функціональне тестування", "FN-20", "Медіа /media/public", "Відображення фото товару",
     "Товар з path у product_images",
     "1. Відкрити URL зображення з картки\n2. Перевірити 404 для неіснуючого path",
     "200 і коректний Content-Type; 404 для чужого шляху",
     "Середній"),

    # --- Кросбраузерне ---
    ("Кросбраузерне тестування", "XB-01", "Вітрина", "Chrome (остання версія)",
     "npm run build; production assets",
     "1. Відкрити головну, каталог, товар, checkout у Chrome\n2. Перевірити JS (кошик, фільтри, галерея)",
     "Візуально ідентично макету; без помилок у Console",
     "Високий"),
    ("Кросбраузерне тестування", "XB-02", "Вітрина", "Mozilla Firefox",
     "Те саме середовище",
     "1. Повторити ключові сценарії FN-01–FN-08 у Firefox",
     "Форми, flex/grid, модалки працюють; шрифти та SVG логотип коректні",
     "Високий"),
    ("Кросбраузерне тестування", "XB-03", "Вітрина", "Microsoft Edge",
     "Те саме середовище",
     "1. Перевірити checkout і Nova Poshta typeahead\n2. LiqPay redirect (sandbox)",
     "Немає блокування popup/redirect; cookies сесії зберігаються",
     "Високий"),
    ("Кросбраузерне тестування", "XB-04", "Вітрина", "Safari (macOS/iOS)",
     "Доступ до Safari або BrowserStack",
     "1. Головна + кошик на iOS Safari\n2. Перевірити position:sticky, input zoom",
     "Немає горизонтального скролу; date/tel inputs без поломки layout",
     "Середній"),
    ("Кросбраузерне тестування", "XB-05", "Адмін-панель", "Адмін у різних браузерах",
     "Обліковий запис admin",
     "1. Завантаження зображення товару\n2. Редагування замовлення в Chrome і Firefox",
     "Upload preview; таблиці orders/reviews без накладання колонок",
     "Середній"),
    ("Кросбраузерне тестування", "XB-06", "Auth", "Google OAuth у браузерах",
     "OAuth налаштовано",
     "1. Login Google у Chrome та Firefox",
     "Callback без CORS/cookie помилок",
     "Середній"),

    # --- Адаптивність ---
    ("Тестування адаптивності", "RS-01", "Layout", "Mobile 320–480px",
     "DevTools device mode або реальний телефон",
     "1. Головна, каталог, кошик, checkout на 375px\n2. Burger-меню header",
     "Одна колонка; кнопки ≥44px; текст читабельний без zoom",
     "Високий"),
    ("Тестування адаптивності", "RS-02", "Layout", "Tablet 768–1024px",
     "—",
     "1. Каталог сітка 2–3 колонки\n2. Сторінка товару: галерея + інфо",
     "Без обрізаних зображень; фільтри доступні",
     "Високий"),
    ("Тестування адаптивності", "RS-03", "Layout", "Desktop ≥1280px",
     "—",
     "1. Wide monitor 1920px\n2. Перевірити max-width контейнера",
     "Контент не розтягнутий на всю ширину; footer/header на місці",
     "Середній"),
    ("Тестування адаптивності", "RS-04", "Checkout", "Форма оформлення на мобільному",
     "Кошик з 1+ товаром",
     "1. Прокрутити wizard доставки\n2. Відкрити клавіатуру на полях tel/email",
     "Поля не перекриваються клавіатурою; кнопка «Оформити» видима",
     "Високий"),
    ("Тестування адаптивності", "RS-05", "Зображення", "Responsive images / lazy",
     "—",
     "1. Network throttling Slow 3G\n2. Скрол каталогу",
     "Зображення підвантажуються прогресивно; placeholder без CLS",
     "Середній"),
    ("Тестування адаптивності", "RS-06", "Адмін", "Адмін-панель на планшеті",
     "Admin login",
     "1. /admin/dashboard на 768px",
     "Таблиці з горизонтальним scroll або stacked; навігація доступна",
     "Низький"),

    # --- Продуктивність ---
    ("Тестування продуктивності", "PF-01", "Сторінки", "TTFB головної та каталогу",
     "Production build, OPcache/cache",
     "1. 10 запитів GET / і /catalog\n2. Зафіксувати TTFB (DevTools/ curl -w)",
     "TTFB < 500 ms на локальному XAMPP; < 800 ms на Render (cold start окремо)",
     "Високий"),
    ("Тестування продуктивності", "PF-02", "Сторінки", "Lighthouse Performance",
     "Chrome Lighthouse",
     "1. Audit головної і product page (mobile)\n2. Записати LCP, CLS, INP",
     "Performance ≥ 70 (mobile) або за внутрішнім SLA; CLS < 0.1",
     "Середній"),
    ("Тестування продуктивності", "PF-03", "API", "Nova Poshta endpoints",
     "—",
     "1. /shipping/nova/cities?q=Київ\n2. branches/streets",
     "Відповідь < 2 с; кешування на клієнті/debounce typeahead",
     "Середній"),
    ("Тестування продуктивності", "PF-04", "Assets", "Vite bundle size",
     "npm run build",
     "1. Перевірити розмір public/build/assets/*.js\n2. Перше завантаження з cache disabled",
     "JS main chunk прийнятний (<500 KB gzip орієнтир); code splitting для admin окремо",
     "Середній"),
    ("Тестування продуктивності", "PF-05", "БД", "Каталог з великою вибіркою",
     "≥100 товарів (seed)",
     "1. /catalog?page=1&filters\n2. Перевірити query count (debugbar/telescope)",
     "Немає N+1 на images/categories; пагінація < 1 с",
     "Високий"),
    ("Тестування продуктивності", "PF-06", "Health", "Endpoint /health",
     "—",
     "1. GET /health під навантаженням 50 rps (короткий burst)",
     "200 JSON {\"ok\":true}; без memory leak на worker",
     "Низький"),

    # --- Безпека ---
    ("Безпекове тестування", "SC-01", "Доступ", "Адмін лише для role admin",
     "Користувач role=user",
     "1. GET /admin після login звичайного user\n2. CSRF POST на admin routes",
     "403/redirect; 419 без CSRF token",
     "Високий"),
    ("Безпекове тестування", "SC-02", "Auth", "Brute-force login",
     "—",
     "1. 10+ невдалих login підряд\n2. Перевірити throttle",
     "Rate limit / lockout; загальне повідомлення без user enumeration",
     "Високий"),
    ("Безпекове тестування", "SC-03", "Checkout", "Throttle оформлення",
     "—",
     "1. Масові POST /checkout",
     "Обмеження middleware throttle:checkout; без дублікатів spam orders",
     "Високий"),
    ("Безпекове тестування", "SC-04", "Input", "XSS у пошуку та відгуках",
     "—",
     "1. Ввести <script>alert(1)</script> у search і review\n2. Переглянути сторінку",
     "Вивід екранований (Blade {{ }}); скрипт не виконується",
     "Високий"),
    ("Безпекове тестування", "SC-05", "Input", "SQL-ін'єкція",
     "Доступ до форм пошуку, login, checkout; логи БД/додатку увімкнено",
     "1. Каталог: /catalog?search=' OR '1'='1 -- та \"1; DROP TABLE products;--\"\n"
     "2. Login: email ' OR 1=1 --@test.com, пароль довільний\n"
     "3. Checkout (POST): ін'єкція в phone, address, promo_code\n"
     "4. /search-suggestions?q=1' UNION SELECT null,null--\n"
     "5. Перевірити відповідь і логи — без SQL-помилок у HTML",
     "Запити параметризовані (Eloquent/Query Builder); 422/400 або порожній безпечний результат; "
     "немає dump SQL/стеку; таблиці БД не змінені; вхід без валідних credentials неможливий",
     "Високий"),
    ("Безпекове тестування", "SC-06", "Payments", "LiqPay signature",
     "—",
     "1. POST /payments/liqpay/callback з невалідним signature\n2. Повтор callback",
     "Відхилено 4xx; подвійна оплата не змінює баланс двічі",
     "Високий"),
    ("Безпекове тестування", "SC-07", "Media", "Path traversal /media/public",
     "—",
     "1. Запит ../../../.env через path\n2. Чужий private file",
     "404/403; лише public disk",
     "Високий"),
    ("Безпекове тестування", "SC-08", "Session", "IDOR замовлення thanks",
     "2 користувачі, 2 orders",
     "1. Відкрити /checkout/thanks/{чужий_id}",
     "403 або generic 404; без PII чужого замовлення",
     "Високий"),
    ("Безпекове тестування", "SC-09", "Headers", "Security headers (prod)",
     "Production deploy",
     "1. Перевірити відповіді на X-Frame-Options, CSP (якщо є), HTTPS redirect",
     "Cookies Secure/HttpOnly/SameSite; HSTS на хостингу",
     "Середній"),
    ("Безпекове тестування", "SC-10", "Secrets", ".env не доступний з web",
     "—",
     "1. GET /.env, /storage/logs/laravel.log",
     "404/403 від веб-сервера",
     "Високий"),
    ("Безпекове тестування", "SC-11", "Newsletter", "Spam підписка",
     "—",
     "1. 20 POST /newsletter/subscribe за хвилину",
     "Throttle 15/min; валідація email",
     "Середній"),
]

HEADERS = [
    "Тип тестування",
    "ID",
    "Модуль",
    "Назва тест-кейсу",
    "Передумови",
    "Кроки перевірки",
    "Очікуваний результат",
    "Пріоритет",
    "Статус",
]


def set_cell_shading(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)

    title = doc.add_heading("Brick Shop — план ручного тестування", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(
        "Проєкт: lego-shop3 (Laravel + Vite)\n"
        "Версія документа: 1.0\n"
        "Категорії: функціональне, кросбраузерне, адаптивність, продуктивність, безпека"
    )
    run.font.size = Pt(10)

    doc.add_paragraph()

    table = doc.add_table(rows=1, cols=len(HEADERS))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = table.rows[0].cells
    for i, text in enumerate(HEADERS):
        hdr_cells[i].text = text
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(9)
        set_cell_shading(hdr_cells[i], "D9E2F3")

    current_category = None
    for row in TESTS:
        cat = row[0]
        cells = table.add_row().cells
        values = list(row[1:]) + ["Не виконано"]
        if cat != current_category:
            current_category = cat
            cells[0].text = cat
        else:
            cells[0].text = ""
        for i, val in enumerate(values):
            cells[i + 1].text = val
            for p in cells[i + 1].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.add_run(
        "Примітка: автоматизовані PHPUnit-тести (tests/Feature) покривають частину функціоналу "
        "(каталог, кошик, checkout, auth, admin). Статус у таблиці оновлюйте після проходження "
        "(Не виконано / Пройдено / Провалено)."
    ).font.size = Pt(9)

    out_path = OUTPUT
    import os
    script_dir = os.path.dirname(os.path.abspath(__file__))
    full_path = os.path.join(script_dir, out_path)
    doc.save(full_path)
    print(full_path)


if __name__ == "__main__":
    main()
